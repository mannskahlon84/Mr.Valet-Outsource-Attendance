"""Excel and PDF downloads for the reports, in one consistent Mr. Valet layout.

A report is a title, optional subtitle lines, a table (headers + rows) and optional summary
lines (label, value) shown under the table.
"""
import io
from datetime import datetime, timezone
from typing import Iterable, List, Optional, Sequence, Tuple

import openpyxl
from fastapi.responses import StreamingResponse
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.core.timeutil import QATAR_TZ

BRAND_GOLD = "DBB457"
BRAND_CHARCOAL = "1A1A1A"

Summary = Sequence[Tuple[str, object]]


def qatar_time(value: Optional[datetime], fmt: str = "%Y-%m-%d %H:%M") -> str:
    """Attendance times are stored as UTC; show them as Qatar time."""
    if not value:
        return ""
    if value.tzinfo is None:
        value = value.replace(tzinfo=timezone.utc)
    return value.astimezone(QATAR_TZ).strftime(fmt)


def _cell(value):
    return "" if value is None else value


def excel_response(filename: str, title: str, headers: List[str], rows: Iterable[Sequence],
                   subtitle: Sequence[str] = (), summary: Summary = ()) -> StreamingResponse:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title[:31]

    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14, color=BRAND_CHARCOAL)
    for line in subtitle:
        ws.append([line])
    ws.append([f"Generated {datetime.now(QATAR_TZ).strftime('%Y-%m-%d %H:%M')} (Qatar time)"])
    ws.append([])

    header_row = ws.max_row + 1
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=header_row, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=BRAND_CHARCOAL)
        cell.alignment = Alignment(vertical="center")
    count = 0
    for row in rows:
        ws.append([_cell(v) for v in row])
        count += 1
    if count == 0:
        ws.append(["No records for this selection."])

    if summary:
        ws.append([])
        for label, value in summary:
            ws.append([label, _cell(value)])
            ws.cell(row=ws.max_row, column=1).font = Font(bold=True)

    # Readable column widths
    for col in range(1, len(headers) + 1):
        letter = get_column_letter(col)
        longest = max((len(str(c.value)) for c in ws[letter][header_row - 1:] if c.value is not None), default=8)
        ws.column_dimensions[letter].width = min(max(longest + 2, 10), 45)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}.xlsx"'},
    )


def pdf_response(filename: str, title: str, headers: List[str], rows: Iterable[Sequence],
                 subtitle: Sequence[str] = (), summary: Summary = ()) -> StreamingResponse:
    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm, title=title)
    styles = getSampleStyleSheet()
    small = styles["BodyText"].clone("small", fontSize=8, leading=10)
    head = styles["BodyText"].clone("head", fontSize=8, leading=10, textColor=colors.white, fontName="Helvetica-Bold")

    elements = [
        Paragraph(f"<font color='#{BRAND_GOLD}'>MR. VALET</font> · {title}", styles["Title"]),
    ]
    for line in subtitle:
        elements.append(Paragraph(line, styles["BodyText"]))
    elements.append(Paragraph(f"Generated {datetime.now(QATAR_TZ).strftime('%Y-%m-%d %H:%M')} (Qatar time)", small))
    elements.append(Spacer(1, 6 * mm))

    data = [[Paragraph(str(h), head) for h in headers]]
    for row in rows:
        data.append([Paragraph(str(_cell(v)), small) for v in row])
    if len(data) == 1:
        data.append([Paragraph("No records for this selection.", small)] + [""] * (len(headers) - 1))

    table = Table(data, repeatRows=1, colWidths=[doc.width / len(headers)] * len(headers))
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{BRAND_CHARCOAL}")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F6F4EF")]),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D9D6CE")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    elements.append(table)

    if summary:
        elements.append(Spacer(1, 6 * mm))
        summary_table = Table([[Paragraph(f"<b>{label}</b>", small), Paragraph(str(_cell(value)), small)] for label, value in summary],
                              colWidths=[60 * mm, 80 * mm], hAlign="LEFT")
        summary_table.setStyle(TableStyle([
            ("LINEABOVE", (0, 0), (-1, 0), 1, colors.HexColor(f"#{BRAND_GOLD}")),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
        ]))
        elements.append(summary_table)

    doc.build(elements)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}.pdf"'},
    )


def export_response(fmt: str, filename: str, title: str, headers: List[str], rows: Iterable[Sequence],
                    subtitle: Sequence[str] = (), summary: Summary = ()) -> StreamingResponse:
    rows = list(rows)
    if fmt == "excel":
        return excel_response(filename, title, headers, rows, subtitle, summary)
    return pdf_response(filename, title, headers, rows, subtitle, summary)
